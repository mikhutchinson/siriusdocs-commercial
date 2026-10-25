"""Verify the installed public API without an engine source checkout."""
import io
import sys
import tempfile
from pathlib import Path

import sirius_docs
from sirius_docs import DocumentExtractor, DocumentLimits, SiriusDocsError

assert sirius_docs.__version__ == "0.1.0"
assert Path(sirius_docs.__file__).suffix == ".pyc"
with DocumentExtractor() as client:
    for kind, source, expected in [
        ("markdown", b"# Source heading\n\n**Exact source**", "**Exact source**"),
        ("text", b"Plain source text", "Plain source text"),
        ("csv", b"Name,Value\nExample,42", "42"),
        ("html", b"<h1>Source</h1><p>Exact text</p><script>hidden</script>", "Exact text"),
        ("asciidoc", b"= Source heading", "= Source heading"),
        ("latex", b"\\section{Source heading}", "\\section{Source heading}"),
    ]:
        result = client.extract(source, format=kind)
        assert result.format.value == kind and expected in result.markdown
        if kind in {"asciidoc", "latex"}:
            assert any(w.code == "source_markup" for w in result.warnings)
    try:
        client.extract(b"PK invalid container")
    except SiriusDocsError as error:
        assert error.code == "invalid_document"
    else:
        raise AssertionError("invalid input succeeded")
    assert "recovery" in client.extract(b"recovery", format="text").markdown
    pdf = client.extract(Path(__file__).with_name("example.pdf"))
    assert pdf.format.value == "pdf" and "Public source document 42." in pdf.markdown
    assert [unit.identifier for unit in pdf.units] == ["page-1"]
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    document = Document()
    document.add_paragraph("First paragraph")
    document.add_table(1, 1).cell(0, 0).text = "Source cell"
    document.add_paragraph("Last paragraph")
    stream = io.BytesIO()
    document.save(stream)
    word = client.extract(stream.getvalue(), format="docx")
    assert word.markdown.index("First paragraph") < word.markdown.index("Source cell") < word.markdown.index("Last paragraph")

    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    stream = io.BytesIO()
    workbook.save(stream)
    sheet = client.extract(stream.getvalue(), format="xlsx")
    cells = sheet.units[0].blocks[0].metadata["cells"]
    assert cells[0]["formula"] == "=1+1" and cells[0]["cached_value"] is None

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Slide source"
    slide.notes_slide.notes_text_frame.text = "Speaker source"
    stream = io.BytesIO()
    deck.save(stream)
    presentation = client.extract(stream.getvalue(), format="pptx")
    assert "Slide source" in presentation.markdown and "Speaker source" in presentation.markdown

    unicode_source = "Résumé Ω 🐳"
    text = client.extract(unicode_source.encode("utf-8"), format="text")
    with tempfile.TemporaryDirectory() as temporary:
        output = Path(temporary) / "export"
        text.write(output)
        assert unicode_source in (output / "document.md").read_bytes().decode("utf-8")
with DocumentExtractor(limits=DocumentLimits(max_bytes=2)) as limited:
    try:
        limited.extract(b"too long", format="text")
    except SiriusDocsError as error:
        assert error.code == "resource_limit"
    else:
        raise AssertionError("limit was ignored")
print("SiriusDocs 0.1.0 compiled consumer passed on " + sys.platform)
